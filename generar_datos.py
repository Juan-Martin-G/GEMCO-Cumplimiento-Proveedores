"""
generar_datos.py
================
Precalcula el cumplimiento de proveedores a partir del Excel consolidado y lo
guarda en docs/datos_cumplimiento.json, que consume el panel estático
(HTML/JS puro en GitHub Pages, sin Python en vivo).

USO (cada vez que se actualiza el Excel):

    python generar_datos.py                      # usa el .xlsx más reciente de data/raw/
    python generar_datos.py "data/raw/archivo.xlsx"

Lee los valores CACHEADOS del Excel (openpyxl data_only=True): el archivo debe
estar guardado desde Excel con las fórmulas calculadas.

Antes de escribir el JSON concilia lo recalculado contra los valores de cada
hoja de marca visible. Si algo no cuadra, termina con error y NO escribe el JSON.
"""

from __future__ import annotations

import json
import math
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

# ──────────────────────────────────────────────────────────────────────
# CONFIGURACIÓN
# ──────────────────────────────────────────────────────────────────────

# False → el JSON no incluye las tablas de BO abiertas ni de facturas
# (ni OC, ni productos, ni precios): solo agregados. El panel oculta esas tablas.
INCLUIR_DETALLE = True

CARPETA_EXCEL = Path("data/raw")
SALIDA_JSON = Path("docs/datos_cumplimiento.json")

# Nombres EXACTOS (sin trim). Ojo: existe una hoja oculta "BBDD Facturacion "
# (con espacio final), base antigua de 2025, que NO debe leerse nunca.
HOJA_FACTURACION = "BBDD Facturacion"
HOJA_BO = "BBDD BO abiertas"
HOJA_MONEDA = "Parametros Moneda"

ENCABEZADOS_FACTURACION = [
    "Proveedor", "Fecha de contabilización", "Año", "Mes", "Número de folio", "Comentarios",
    "Orden de compra", "Categoria", "Moneda del documento", "Total documento (ME)", "Moneda marca",
    "TC a moneda marca", "Total moneda marca",
]
ENCABEZADOS_BO = [
    "Fecha OC", "Año OC", "Mes", "Referencia", "Producto", "Descripcion", "Grupo", "Marca",
    "Cantidad pdte", "moneda", "Precio", "Total Neto ME", "Moneda marca", "TC a moneda marca",
    "Total moneda marca",
]

CATEGORIAS = ["Equipos", "Repuestos", "Accesorios", "Consumibles"]
SIN_CATEGORIA = "Sin categoría"

TOLERANCIA_MONTO = 0.01
TOLERANCIA_PCT = 0.0001

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
MESES_CORTOS = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

# Desplazamientos relativos a la columna de etiquetas ("Item") de cada bloque.
LAYOUT_MESES = [1, 2, 3, 5, 6, 7, 9, 10, 11, 13, 14, 15]
LAYOUT_TRIMESTRES = [4, 8, 12, 16]
LAYOUT_TOTAL = 17
FILA = {"fact": 3, "bo": 4, "factbo": 5, "forecast": 6, "cumpl": 7, "bo_total": 15, "fact_total": 23}


class ErrorCarga(Exception):
    """Error que detiene la carga (hoja o encabezados faltantes, etc.)."""


# ──────────────────────────────────────────────────────────────────────
# UTILIDADES
# ──────────────────────────────────────────────────────────────────────

NO_NUM = object()  # celda con valor no numérico (texto, #N/A, etc.)


def txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def clave(v) -> str:
    return txt(v).lower()


def numero(v):
    """None = vacía · NO_NUM = no numérica · float en otro caso."""
    if v is None:
        return None
    if isinstance(v, bool):
        return NO_NUM
    if isinstance(v, (int, float)):
        return NO_NUM if isinstance(v, float) and math.isnan(v) else float(v)
    if isinstance(v, str) and v.strip() == "":
        return None
    return NO_NUM


def es_num(x) -> bool:
    return x is not None and x is not NO_NUM


def a_fecha(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return from_excel(v)
        except (ValueError, OverflowError):
            return None
    return None


def idx_mes(anio: int, mes: int) -> int:
    return anio * 12 + mes


def mes_corto(anio: int, mes: int) -> str:
    return f"{MESES_CORTOS[mes]}-{str(anio)[2:]}"


def col_letra(c0: int) -> str:
    return get_column_letter(c0 + 1)


def ref_hoja(h: str) -> str:
    return h if h.replace("_", "").isalnum() and h.isascii() else "'" + h.replace("'", "''") + "'"


def dir_celda(fila: int, c0: int) -> str:
    return f"{col_letra(c0)}{fila}"


def valor(ws, fila: int, c0: int):
    """fila 1-based, columna 0-based."""
    return ws.cell(row=fila, column=c0 + 1).value


def categoria(v) -> str:
    k = clave(v)
    for c in CATEGORIAS:
        if c.lower() == k:
            return c
    return SIN_CATEGORIA


def todas_categorias():
    return CATEGORIAS + [SIN_CATEGORIA]


def vacio_cat():
    return {c: 0.0 for c in todas_categorias()}


def hoja_oculta(ws) -> bool:
    return ws.sheet_state != "visible"


# ──────────────────────────────────────────────────────────────────────
# LECTURA DE BBDD (por nombre de encabezado)
# ──────────────────────────────────────────────────────────────────────

def leer_tabla(wb, nombre: str, requeridos: list[str]) -> dict:
    # Búsqueda por nombre EXACTO: "BBDD Facturacion" ≠ "BBDD Facturacion ".
    if nombre not in wb.sheetnames:
        hojas = ", ".join(f'"{n}"' for n in wb.sheetnames)
        raise ErrorCarga(f'No se encontró la hoja "{nombre}" (se busca por nombre exacto). Hojas del libro: {hojas}.')
    ws = wb[nombre]
    filas = ws.iter_rows(min_row=1, min_col=1, values_only=True)
    encabezado = next(filas, None)
    if not encabezado:
        raise ErrorCarga(f'La hoja "{nombre}" está vacía.')
    col: dict[str, int] = {}
    for i, h in enumerate(encabezado):
        k = txt(h)
        if k and k not in col:
            col[k] = i
    faltan = [h for h in requeridos if h not in col]
    if faltan:
        raise ErrorCarga(f'La hoja "{nombre}" no tiene los encabezados requeridos: '
                         + ", ".join(f'"{h}"' for h in faltan)
                         + ". Se detuvo la carga para no leer una base equivocada.")
    datos = []
    for n, fila in enumerate(filas, start=2):
        valores = {h: (fila[col[h]] if col[h] < len(fila) else None) for h in requeridos}
        if all(txt(v) == "" for v in valores.values()):
            continue
        datos.append({"fila": n, "c": valores})
    return {"nombre": nombre, "filas": datos, "oculta": hoja_oculta(ws)}


MAPEO = {
    "fact": {"marca": "Proveedor", "fecha": "Fecha de contabilización", "anio": "Año", "mes": "Mes",
             "moneda_orig": "Moneda del documento", "monto_orig": "Total documento (ME)", "grupo": "Categoria"},
    "bo": {"marca": "Marca", "fecha": "Fecha OC", "anio": "Año OC", "mes": "Mes",
           "moneda_orig": "moneda", "monto_orig": "Total Neto ME", "grupo": "Grupo"},
}


def normalizar(tabla: dict, tipo: str, chk: dict) -> dict:
    m = MAPEO[tipo]
    registros = []
    corte = None
    for f in tabla["filas"]:
        c = f["c"]
        ref = f'{tabla["nombre"]} fila {f["fila"]}'
        fecha = a_fecha(c[m["fecha"]])
        if fecha is None:
            chk["fechas"].append(ref)
            continue
        if corte is None or fecha > corte:
            corte = fecha
        anio, mes = fecha.year, fecha.month - 1
        if numero(c[m["anio"]]) != anio or clave(c[m["mes"]]) != MESES[mes]:
            chk["anio_mes"].append(ref)

        total = numero(c["Total moneda marca"])
        if total is NO_NUM:
            chk["totales"].append(ref)
        tc = numero(c["TC a moneda marca"])
        if not es_num(tc):
            chk["tc"].append(f'{ref} ("{txt(c["TC a moneda marca"])}")')
        monto_orig = numero(c[m["monto_orig"]])

        reg = {
            "tipo": tipo, "fila": f["fila"], "fecha": fecha, "anio": anio, "mes": mes,
            "marca": txt(c[m["marca"]]), "marca_clave": clave(c[m["marca"]]),
            "moneda_orig": txt(c[m["moneda_orig"]]),
            "monto_orig": monto_orig if es_num(monto_orig) else 0.0,
            "moneda_marca": txt(c["Moneda marca"]),
            "tc": tc if es_num(tc) else None,
            "total": total if es_num(total) else 0.0,
            "cat_raw": txt(c[m["grupo"]]),
            "cat": categoria(c[m["grupo"]]),
        }
        if tipo == "fact":
            reg["folio"] = txt(c["Número de folio"])
            reg["comentarios"] = txt(c["Comentarios"])
            reg["oc"] = txt(c["Orden de compra"])
        else:
            cant = numero(c["Cantidad pdte"])
            precio = numero(c["Precio"])
            reg["referencia"] = txt(c["Referencia"])
            reg["producto"] = txt(c["Producto"])
            reg["descripcion"] = txt(c["Descripcion"])
            reg["cantidad"] = cant if es_num(cant) else 0.0
            reg["precio"] = precio if es_num(precio) else 0.0
        reg["convertido"] = bool(reg["moneda_orig"]) and clave(reg["moneda_orig"]) != clave(reg["moneda_marca"])
        registros.append(reg)
    return {"registros": registros, "corte": corte}


def indexar(registros: list[dict]) -> dict:
    """(marca, año, mes) → {total, cat, filas}"""
    idx: dict = {}
    for r in registros:
        k = (r["marca_clave"], r["anio"], r["mes"])
        e = idx.get(k)
        if e is None:
            e = idx[k] = {"total": 0.0, "cat": vacio_cat(), "filas": []}
        e["total"] += r["total"]
        e["cat"][r["cat"]] += r["total"]
        e["filas"].append(r)
    return idx


def leer_parametros_moneda(wb) -> dict:
    if HOJA_MONEDA not in wb.sheetnames:
        raise ErrorCarga(f'No se encontró la hoja "{HOJA_MONEDA}".')
    ws = wb[HOJA_MONEDA]
    marcas, pares = [], {}
    for r in range(2, ws.max_row + 1):
        marca, mon = valor(ws, r, 0), valor(ws, r, 1)
        if txt(marca):
            marcas.append({"marca": str(marca), "clave": clave(marca), "moneda": txt(mon), "celda": f"A{r}"})
        par, tc = txt(valor(ws, r, 3)), numero(valor(ws, r, 4))
        if par:
            pares[par] = {"tc": tc if es_num(tc) else None, "celda": f"E{r}"}
    return {"marcas": marcas, "pares": pares}


# ──────────────────────────────────────────────────────────────────────
# HOJAS DE MARCA
# ──────────────────────────────────────────────────────────────────────

def leer_bloque(ws, hoja: str, c0: int, tipo: str) -> dict:
    errores = []
    meses = []
    for off in LAYOUT_MESES:
        c = c0 + off
        etiqueta = txt(valor(ws, 2, c))
        mes = MESES.index(etiqueta.lower()) if etiqueta.lower() in MESES else -1
        anio = numero(valor(ws, 1, c))
        if mes < 0:
            errores.append(f'{hoja}!{dir_celda(2, c)}: se esperaba un mes y dice "{etiqueta}".')
        if not (es_num(anio) and anio > 0):
            errores.append(f"{hoja}!{dir_celda(1, c)}: falta el año del mes.")
            anio = None
        meses.append({"col": c, "mes": mes, "anio": int(anio) if anio else None})
    for q, off in enumerate(LAYOUT_TRIMESTRES):
        e = clave(valor(ws, 2, c0 + off))
        if e != f"q{q + 1}":
            errores.append(f'{hoja}!{dir_celda(2, c0 + off)}: se esperaba "Q{q + 1}" y dice "{e}".')
    if clave(valor(ws, 2, c0 + LAYOUT_TOTAL)) != "total":
        errores.append(f'{hoja}!{dir_celda(2, c0 + LAYOUT_TOTAL)}: se esperaba "Total".')
    if not errores:
        for i in range(1, 12):
            if idx_mes(meses[i]["anio"], meses[i]["mes"]) != idx_mes(meses[i - 1]["anio"], meses[i - 1]["mes"]) + 1:
                errores.append(f"{hoja}: los meses del bloque no son consecutivos ({dir_celda(2, meses[i]['col'])}).")
                break

    # Forecast (fila 6): celdas de mes, trimestre y total tal como están en la hoja.
    def fc_celda(off):
        n = numero(valor(ws, FILA["forecast"], c0 + off))
        return n if es_num(n) else None

    fc_mes = [fc_celda(off) for off in LAYOUT_MESES]
    fc_trim = [fc_celda(off) for off in LAYOUT_TRIMESTRES]
    fc_total = fc_celda(LAYOUT_TOTAL)
    fc_derivado = [False] * 12
    for t, q in enumerate(fc_trim):
        ids = [3 * t, 3 * t + 1, 3 * t + 2]
        if q and q > 0 and all(not fc_mes[i] for i in ids):
            for i in ids:
                fc_mes[i] = q / 3
                fc_derivado[i] = True

    ok = not errores
    ini, fin = meses[0], meses[11]
    if tipo == "calendario":
        titulo = f"Año calendario {ini['anio']}" if ok else "Año calendario"
    else:
        titulo = (f"Año fiscal USA ({mes_corto(ini['anio'], ini['mes'])} – {mes_corto(fin['anio'], fin['mes'])})"
                  if ok else "Año fiscal USA")
    return {"tipo": tipo, "hoja": hoja, "c0": c0, "titulo": titulo, "meses": meses, "errores": errores,
            "fc_mes": fc_mes, "fc_trim": fc_trim, "fc_total": fc_total, "fc_derivado": fc_derivado, "ws": ws}


def calcular_bloque(b: dict, marca_clave: str, idx_f: dict, idx_b: dict, corte_idx: int) -> dict:
    vacio = {"total": 0.0, "cat": vacio_cat(), "filas": []}
    meses = []
    for i, m in enumerate(b["meses"]):
        F = idx_f.get((marca_clave, m["anio"], m["mes"]), vacio)
        B = idx_b.get((marca_clave, m["anio"], m["mes"]), vacio)
        meses.append({
            "anio": m["anio"], "mes": m["mes"], "etiqueta": mes_corto(m["anio"], m["mes"]),
            "fact": F["total"], "bo": B["total"], "factbo": F["total"] + B["total"],
            "fc": b["fc_mes"][i], "fcDerivado": b["fc_derivado"][i],
            "_catF": F["cat"], "_catB": B["cat"], "_filasF": F["filas"], "_filasB": B["filas"],
        })

    def suma(arr, k):
        return sum((x[k] or 0) for x in arr)

    trims = []
    for q in range(4):
        ms = meses[3 * q:3 * q + 3]
        fc = b["fc_trim"][q]
        factbo = suma(ms, "factbo")
        primero, ultimo = idx_mes(ms[0]["anio"], ms[0]["mes"]), idx_mes(ms[2]["anio"], ms[2]["mes"])
        estado = "cerrado" if corte_idx >= ultimo else "parcial" if corte_idx >= primero else "futuro"
        trims.append({"nombre": f"Q{q + 1}", "fact": suma(ms, "fact"), "bo": suma(ms, "bo"), "factbo": factbo,
                      "fc": fc, "cumpl": factbo / fc if fc and fc > 0 else None, "estado": estado})

    total = {"fact": suma(trims, "fact"), "bo": suma(trims, "bo"), "factbo": suma(trims, "factbo"), "fc": b["fc_total"]}
    total["cumpl"] = total["factbo"] / total["fc"] if total["fc"] and total["fc"] > 0 else None

    ms_ytd = [m for m in meses if idx_mes(m["anio"], m["mes"]) <= corte_idx]
    ytd = {"fact": suma(ms_ytd, "fact"), "bo": suma(ms_ytd, "bo"), "factbo": suma(ms_ytd, "factbo"),
           "fc": suma(ms_ytd, "fc"), "nMeses": len(ms_ytd),
           "desde": {"anio": meses[0]["anio"], "mes": meses[0]["mes"]},
           "hasta": {"anio": ms_ytd[-1]["anio"], "mes": ms_ytd[-1]["mes"]} if ms_ytd else None,
           "completo": len(ms_ytd) == 12}
    ytd["cumpl"] = ytd["factbo"] / ytd["fc"] if ytd["fc"] > 0 else None

    cat_f, cat_b, raw_f, raw_b = vacio_cat(), vacio_cat(), {}, {}
    for m in meses:
        for c in todas_categorias():
            cat_f[c] += m["_catF"][c]
            cat_b[c] += m["_catB"][c]
        for r in m["_filasF"]:
            if r["cat"] == SIN_CATEGORIA:
                k = r["cat_raw"] or "(vacío)"
                raw_f[k] = raw_f.get(k, 0.0) + r["total"]
        for r in m["_filasB"]:
            if r["cat"] == SIN_CATEGORIA:
                k = r["cat_raw"] or "(vacío)"
                raw_b[k] = raw_b.get(k, 0.0) + r["total"]

    return {
        "meses": meses, "trims": trims, "total": total, "ytd": ytd,
        "catF": cat_f, "catB": cat_b, "rawF": raw_f, "rawB": raw_b,
        "filasF": [r for m in meses for r in m["_filasF"]],
        "filasB": [r for m in meses for r in m["_filasB"]],
        "fcSumaMeses": suma(meses, "fc"),
        "tieneForecast": bool(total["fc"] and total["fc"] > 0),
        "hayDerivado": any(b["fc_derivado"]),
    }


def conciliar_bloque(b: dict, calc: dict) -> list[dict]:
    ws, c0, hoja = b["ws"], b["c0"], b["hoja"]
    items = []

    def add(concepto, fila, off, app, tipo, nota=""):
        ex = numero(valor(ws, fila, c0 + off))
        no_num = ex is NO_NUM
        ex_v = 0.0 if ex is None else (None if no_num else ex)  # vacío ≙ 0 (IFERROR del Excel)
        ap_v = 0.0 if app is None else app
        dif = None if ex_v is None else ap_v - ex_v
        tol = TOLERANCIA_PCT if tipo == "pct" else TOLERANCIA_MONTO
        ok = dif is not None and abs(dif) <= tol + 1e-9
        items.append({"concepto": concepto, "celda": f"{ref_hoja(hoja)}!{dir_celda(fila, c0 + off)}",
                      "excel": None if (ex is None or no_num) else ex, "excelVacia": ex is None,
                      "excelNoNumerico": no_num, "app": app, "dif": dif, "ok": ok, "tipo": tipo, "nota": nota})

    def sum4(cat):
        return sum(cat[c] for c in CATEGORIAS)

    T = LAYOUT_TOTAL
    add("Facturación total (fila 3)", FILA["fact"], T, calc["total"]["fact"], "monto")
    add("Backorder total (fila 4)", FILA["bo"], T, calc["total"]["bo"], "monto")
    add("Fact + BO total (fila 5)", FILA["factbo"], T, calc["total"]["factbo"], "monto")
    add("Forecast total (fila 6)", FILA["forecast"], T, calc["fcSumaMeses"], "monto",
        "Σ forecast mensual usado (incluye meses derivados = trimestre/3)" if calc["hayDerivado"] else "Σ forecast mensual usado")
    add("Cumplimiento total (fila 7)", FILA["cumpl"], T, calc["total"]["cumpl"], "pct",
        'Sin forecast: el panel muestra "—"; el Excel muestra 0 por IFERROR' if calc["total"]["cumpl"] is None else "")
    for q, t in enumerate(calc["trims"]):
        add(f'Cumplimiento {t["nombre"]} (fila 7)', FILA["cumpl"], LAYOUT_TRIMESTRES[q], t["cumpl"], "pct",
            "Sin forecast en el trimestre (Excel: 0 por IFERROR)" if t["cumpl"] is None else "")
    sin_b, sin_f = calc["catB"][SIN_CATEGORIA], calc["catF"][SIN_CATEGORIA]
    add("Total BO por categoría (fila 15)", FILA["bo_total"], T, sum4(calc["catB"]), "monto",
        f'Excluye "{SIN_CATEGORIA}" ({sin_b}), igual que la hoja' if sin_b else "")
    add("Total facturación por categoría (fila 23)", FILA["fact_total"], T, sum4(calc["catF"]), "monto",
        f'Excluye "{SIN_CATEGORIA}" ({sin_f}), igual que la hoja' if sin_f else "")
    return items


def resumen_conversion(calc: dict) -> list[dict]:
    """Monto original, TC y monto convertido por moneda original (recuadro "Conversión de moneda")."""
    salida = []
    for base, filas in ((HOJA_FACTURACION, calc["filasF"]),
                        (HOJA_BO, [r for r in calc["filasB"] if r["cantidad"] > 0 or r["total"]])):
        g: dict = {}
        for r in filas:
            k = r["moneda_orig"] or "(vacía)"
            e = g.get(k)
            if e is None:
                e = g[k] = {"base": base, "moneda": k, "n": 0, "orig": 0.0, "conv": 0.0, "tcs": [], "convertido": r["convertido"]}
            e["n"] += 1
            e["orig"] += r["monto_orig"]
            e["conv"] += r["total"]
            if r["tc"] not in e["tcs"]:
                e["tcs"].append(r["tc"])
        salida.extend(g.values())
    return salida


# ──────────────────────────────────────────────────────────────────────
# PROCESO COMPLETO
# ──────────────────────────────────────────────────────────────────────

def _lista(arr, n=8):
    return " · ".join(arr[:n]) + (f" · … ({len(arr)} en total)" if len(arr) > n else "")


def _fila_detalle_fact(r):
    return {"fecha": r["fecha"].strftime("%Y-%m-%d"), "anio": r["anio"], "mes": r["mes"], "folio": r["folio"],
            "comentarios": r["comentarios"], "oc": r["oc"], "cat": r["cat"], "catRaw": r["cat_raw"],
            "monedaOrig": r["moneda_orig"], "montoOrig": r["monto_orig"], "tc": r["tc"], "total": r["total"],
            "convertido": r["convertido"]}


def _fila_detalle_bo(r):
    return {"fecha": r["fecha"].strftime("%Y-%m-%d"), "anio": r["anio"], "mes": r["mes"], "referencia": r["referencia"],
            "producto": r["producto"], "descripcion": r["descripcion"], "cat": r["cat"], "cantidad": r["cantidad"],
            "monedaOrig": r["moneda_orig"], "precio": r["precio"], "montoOrig": r["monto_orig"], "tc": r["tc"],
            "total": r["total"], "convertido": r["convertido"]}


def procesar_libro(ruta: Path | str, incluir_detalle: bool = INCLUIR_DETALLE) -> dict:
    """Lee el Excel, calcula y concilia. Devuelve el diccionario que se serializa a JSON."""
    ruta = Path(ruta)
    wb = openpyxl.load_workbook(ruta, data_only=True)
    problemas = []
    chk = {"fechas": [], "anio_mes": [], "totales": [], "tc": []}

    t_f = leer_tabla(wb, HOJA_FACTURACION, ENCABEZADOS_FACTURACION)
    t_b = leer_tabla(wb, HOJA_BO, ENCABEZADOS_BO)
    params = leer_parametros_moneda(wb)
    n_f = normalizar(t_f, "fact", chk)
    n_b = normalizar(t_b, "bo", chk)
    idx_f, idx_b = indexar(n_f["registros"]), indexar(n_b["registros"])
    cortes = [c for c in (n_f["corte"], n_b["corte"]) if c]
    corte = max(cortes) if cortes else None
    corte_idx = idx_mes(corte.year, corte.month - 1) if corte else 10 ** 9

    problemas.append({"nivel": "ok", "titulo": f'Se leyó la hoja "{t_f["nombre"]}" (nombre exacto) · {len(t_f["filas"])} filas · encabezados validados'})
    problemas.append({"nivel": "ok", "titulo": f'Se leyó la hoja "{t_b["nombre"]}" (nombre exacto) · {len(t_b["filas"])} filas · encabezados validados'})
    for n in wb.sheetnames:
        if n != t_f["nombre"] and n.strip() == t_f["nombre"].strip():
            oculta = " (oculta)" if hoja_oculta(wb[n]) else ""
            problemas.append({"nivel": "info", "titulo": f'Se ignoró la hoja "{n.replace(" ", "␣")}"{oculta}: nombre casi igual a "{t_f["nombre"]}", es la base antigua y no se lee.'})
    if t_f["oculta"]:
        problemas.append({"nivel": "aviso", "titulo": f'La hoja "{t_f["nombre"]}" está oculta; se esperaba visible.'})

    if chk["fechas"]:
        problemas.append({"nivel": "error", "bloqueante": True, "titulo": f'{len(chk["fechas"])} filas sin fecha válida (no se suman)', "detalle": _lista(chk["fechas"])})
    if chk["totales"]:
        problemas.append({"nivel": "error", "bloqueante": True, "titulo": f'{len(chk["totales"])} filas con "Total moneda marca" no numérico', "detalle": _lista(chk["totales"])})
    if chk["tc"]:
        problemas.append({"nivel": "error", "bloqueante": True, "titulo": f'{len(chk["tc"])} filas con "TC a moneda marca" no numérico', "detalle": _lista(chk["tc"])})
    else:
        problemas.append({"nivel": "ok", "titulo": 'Todos los "TC a moneda marca" son numéricos'})
    if chk["anio_mes"]:
        problemas.append({"nivel": "aviso", "titulo": f'{len(chk["anio_mes"])} filas donde Año/Mes cacheados no coinciden con la fecha (se usa la fecha)', "detalle": _lista(chk["anio_mes"])})
    else:
        problemas.append({"nivel": "ok", "titulo": "Año y Mes derivados de la fecha coinciden con las columnas del Excel en todas las filas"})

    ocultas, marcas = [], []
    for nombre in wb.sheetnames:
        ws = wb[nombre]
        if hoja_oculta(ws):
            ocultas.append(nombre)
            continue
        if clave(valor(ws, 2, 0)) != "item" or clave(valor(ws, 6, 0)) != "forecast":
            continue
        nombre_marca = txt(valor(ws, 1, 0))
        mc = clave(nombre_marca)
        bloques = [leer_bloque(ws, nombre, 0, "calendario")]
        for c in range(19, ws.max_column):
            if clave(valor(ws, 2, c)) == "item":
                bloques.append(leer_bloque(ws, nombre, c, "fiscal"))
                break
        for b in bloques:
            if b["errores"]:
                problemas.append({"nivel": "error", "bloqueante": True, "titulo": f'Hoja "{nombre}", bloque {b["tipo"]}: estructura inesperada', "detalle": " ".join(b["errores"])})
                b["calc"], b["conc"], b["conversion"] = None, [], []
                continue
            b["calc"] = calcular_bloque(b, mc, idx_f, idx_b, corte_idx)
            b["conc"] = conciliar_bloque(b, b["calc"])
            b["conversion"] = resumen_conversion(b["calc"])

        # Moneda marca única en ambas BBDD
        regs = [r for r in n_f["registros"] if r["marca_clave"] == mc] + [r for r in n_b["registros"] if r["marca_clave"] == mc]
        por_moneda: dict = {}
        for r in regs:
            k = r["moneda_marca"] if r["moneda_marca"] and r["moneda_marca"] != "0" else "(vacía o 0)"
            por_moneda.setdefault(k, []).append(r)
        monedas = sorted(por_moneda, key=lambda k: -len(por_moneda[k]))
        param = next((p for p in params["marcas"] if p["clave"] == mc), None)
        moneda = next((m for m in monedas if m != "(vacía o 0)"), None) or (param and param["moneda"]) or "—"
        if len(monedas) > 1:
            def desc(k):
                rs = por_moneda[k]
                s = f"{k}: {len(rs)} filas"
                if len(rs) <= 3:
                    s += " (" + "; ".join(f'{HOJA_FACTURACION if r["tipo"] == "fact" else HOJA_BO} fila {r["fila"]}, total {r["total"]:g}' for r in rs) + ")"
                return s
            problemas.append({"nivel": "error", "marca": nombre_marca,
                              "titulo": f'{nombre_marca}: tiene {len(monedas)} valores distintos de "Moneda marca" en las BBDD',
                              "detalle": " · ".join(desc(k) for k in monedas)})

        # Conversión de moneda: TC usados vs Parametros Moneda
        convertidos = [r for r in regs if r["convertido"]]
        distintos = []
        for r in convertidos:
            par = params["pares"].get(f'{r["moneda_orig"]}>{r["moneda_marca"]}')
            if not par or par["tc"] != r["tc"]:
                distintos.append(r)
        if distintos:
            problemas.append({"nivel": "aviso", "marca": nombre_marca,
                              "titulo": f'{nombre_marca}: {len(distintos)} filas convertidas con TC distinto al de "{HOJA_MONEDA}"',
                              "detalle": " · ".join(f'fila {r["fila"]}: {r["moneda_orig"]}>{r["moneda_marca"]} TC {r["tc"]}' for r in distintos[:6])})
        pares_conv = []
        for r in convertidos:
            p = {"orig": r["moneda_orig"], "dest": r["moneda_marca"], "tc": r["tc"]}
            if p not in pares_conv:
                pares_conv.append(p)

        marca = {"hoja": nombre, "nombre": nombre_marca, "clave": mc, "moneda": moneda, "monedas": monedas,
                 "bloques": bloques, "tieneFiscal": len(bloques) > 1, "convierte": bool(convertidos),
                 "paresConversion": pares_conv, "sinFilas": not regs}
        if incluir_detalle:
            periodos = {(m["anio"], m["mes"]) for b in bloques for m in b["meses"]}
            marca["detalle"] = {
                "facturas": [_fila_detalle_fact(r) for r in regs if r["tipo"] == "fact" and (r["anio"], r["mes"]) in periodos],
                "bo": [_fila_detalle_bo(r) for r in regs if r["tipo"] == "bo" and r["cantidad"] > 0],
            }
        marcas.append(marca)

    if not marcas:
        raise ErrorCarga('No se encontraron hojas de marca visibles (A2 = "Item" y A6 = "Forecast").')
    ok_m = [m["nombre"] for m in marcas if len(m["monedas"]) <= 1]
    if ok_m:
        problemas.append({"nivel": "ok", "titulo": "Moneda marca única en ambas BBDD: " + ", ".join(ok_m)})
    for m in marcas:
        if m["sinFilas"]:
            problemas.append({"nivel": "info", "titulo": f'{m["nombre"]}: no tiene filas en ninguna BBDD'})
    if ocultas:
        problemas.append({"nivel": "info", "titulo": "Hojas ocultas ignoradas (no se leen): " + ", ".join(f'"{n}"' for n in ocultas)})

    items = [i for m in marcas for b in m["bloques"] for i in b["conc"]]
    resumen = {
        "comparaciones": len(items),
        "ok": sum(1 for i in items if i["ok"]),
        "diferencias": sum(1 for i in items if not i["ok"]),
        "errores": sum(1 for p in problemas if p["nivel"] == "error"),
        "erroresBloqueantes": sum(1 for p in problemas if p.get("bloqueante")),
        "avisos": sum(1 for p in problemas if p["nivel"] == "aviso"),
    }

    def fecha_txt(d):
        return d.strftime("%Y-%m-%d") if d else None

    return {
        "meta": {
            "archivo": ruta.name,
            "generado": datetime.now().isoformat(timespec="seconds"),
            "incluirDetalle": incluir_detalle,
            "hojas": {"facturacion": HOJA_FACTURACION, "bo": HOJA_BO, "moneda": HOJA_MONEDA},
            "fact": {"hoja": t_f["nombre"], "filas": len(t_f["filas"]), "corte": fecha_txt(n_f["corte"])},
            "bo": {"hoja": t_b["nombre"], "filas": len(t_b["filas"]), "corte": fecha_txt(n_b["corte"])},
            "corte": fecha_txt(corte),
            "tolerancia": {"monto": TOLERANCIA_MONTO, "pct": TOLERANCIA_PCT},
            "categorias": CATEGORIAS, "sinCategoria": SIN_CATEGORIA,
        },
        "resumen": resumen,
        "problemas": problemas,
        "ocultas": ocultas,
        "parametros": {"pares": params["pares"], "marcas": [{"marca": p["marca"], "moneda": p["moneda"]} for p in params["marcas"]]},
        "marcas": [_serializar_marca(m) for m in marcas],
    }


def _serializar_marca(m: dict) -> dict:
    salida = {k: v for k, v in m.items() if k != "bloques"}
    salida["bloques"] = []
    for b in m["bloques"]:
        sb = {"tipo": b["tipo"], "titulo": b["titulo"], "c0": b["c0"], "errores": b["errores"],
              "meses": [{"anio": x["anio"], "mes": x["mes"]} for x in b["meses"]],
              "conc": b["conc"], "conversion": b["conversion"], "calc": None}
        c = b["calc"]
        if c:
            sb["calc"] = {
                "meses": [{k: v for k, v in x.items() if not k.startswith("_")} for x in c["meses"]],
                **{k: c[k] for k in ("trims", "total", "ytd", "catF", "catB", "rawF", "rawB", "fcSumaMeses", "tieneForecast", "hayDerivado")},
            }
        salida["bloques"].append(sb)
    return salida


def _redondear(v):
    if isinstance(v, float):
        return round(v, 6)
    if isinstance(v, dict):
        return {k: _redondear(x) for k, x in v.items()}
    if isinstance(v, list):
        return [_redondear(x) for x in v]
    return v


def buscar_excel(carpeta: Path = CARPETA_EXCEL) -> Path:
    candidatos = [p for p in carpeta.glob("*.xlsx") if not p.name.startswith("~$")]
    if not candidatos:
        raise ErrorCarga(f"No hay ningún .xlsx en {carpeta}/. Copia ahí el Excel consolidado.")
    return max(candidatos, key=lambda p: p.stat().st_mtime)


def fmt_cl(v: float, decimales: int = 2) -> str:
    """Formato es-CL: miles con punto, decimales con coma."""
    return f"{v:,.{decimales}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def imprimir_conciliacion(datos: dict) -> None:
    for m in datos["marcas"]:
        for b in m["bloques"]:
            print(f'\n## {m["nombre"]} ({m["moneda"]}) · {b["titulo"]}')
            for e in b["errores"]:
                print(f"  ✕ {e}")
            for it in b["conc"]:
                fmt = (lambda v: "—" if v is None else fmt_cl(v * 100, 1) + "%") if it["tipo"] == "pct" else (lambda v: "—" if v is None else fmt_cl(v))
                ex = "no numérica" if it["excelNoNumerico"] else fmt(it["excel"] if it["excel"] is not None else 0.0)
                print(f'  {"✔" if it["ok"] else "✕"} {it["concepto"]:<42} {it["celda"]:<18} excel={ex:>15} app={fmt(it["app"]):>15}'
                      + ("" if it["ok"] else f'   dif={it["dif"]}'))


def main(argv: list[str]) -> int:
    try:
        ruta = Path(argv[1]) if len(argv) > 1 else buscar_excel()
        print(f"Leyendo {ruta} …")
        datos = procesar_libro(ruta, INCLUIR_DETALLE)
    except ErrorCarga as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1

    meta, r = datos["meta"], datos["resumen"]
    print(f'Facturación: hoja "{meta["fact"]["hoja"]}" · {meta["fact"]["filas"]} filas · corte {meta["fact"]["corte"]}')
    print(f'Backorder:   hoja "{meta["bo"]["hoja"]}" · {meta["bo"]["filas"]} filas · corte {meta["bo"]["corte"]}')
    print("\nValidaciones:")
    for p in datos["problemas"]:
        print(f'  [{p["nivel"]}] {p["titulo"]}' + (f'\n           {p["detalle"]}' if p.get("detalle") else ""))
    imprimir_conciliacion(datos)
    print(f'\nConciliación: {r["ok"]}/{r["comparaciones"]} cuadran · {r["diferencias"]} diferencias · '
          f'{r["errores"]} errores de datos ({r["erroresBloqueantes"]} bloqueantes) · {r["avisos"]} avisos')

    if r["diferencias"] or r["erroresBloqueantes"]:
        print(f"\nERROR: la conciliación no cuadra. NO se escribió {SALIDA_JSON}.", file=sys.stderr)
        return 1

    SALIDA_JSON.parent.mkdir(parents=True, exist_ok=True)
    with open(SALIDA_JSON, "w", encoding="utf-8") as f:
        json.dump(_redondear(datos), f, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    print(f"\nOK: {SALIDA_JSON} ({SALIDA_JSON.stat().st_size / 1024:,.0f} KB, detalle {'incluido' if INCLUIR_DETALLE else 'omitido'}).")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
