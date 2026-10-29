"""
Tests de generar_datos.py

- Regresión con los valores de referencia validados contra el Excel (archivo 20261026).
  Requieren el Excel en data/raw/; si no está, se omiten.
- Reglas de negocio con libros sintéticos (no requieren el Excel).
- Consistencia del JSON publicado en docs/.

    pytest tests/ -v
"""

import json
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import generar_datos as gd

RAIZ = Path(__file__).resolve().parent.parent


# ──────────────────────────────────────────────────────────────────────
# Fixtures
# ──────────────────────────────────────────────────────────────────────

def _excel_real():
    candidatos = [p for p in (RAIZ / "data" / "raw").glob("*.xlsx") if not p.name.startswith("~$")]
    return max(candidatos, key=lambda p: p.stat().st_mtime) if candidatos else None


@pytest.fixture(scope="module")
def datos():
    ruta = _excel_real()
    if ruta is None:
        pytest.skip("No hay Excel en data/raw/ (los tests de referencia necesitan el archivo real).")
    return gd.procesar_libro(ruta, incluir_detalle=True)


def marca(datos, nombre):
    return next(m for m in datos["marcas"] if m["nombre"] == nombre)


def bloque(datos, nombre, tipo="calendario"):
    return next(b for b in marca(datos, nombre)["bloques"] if b["tipo"] == tipo)["calc"]


def pct(v, decimales=0):
    return None if v is None else round(v * 100, decimales)


# ──────────────────────────────────────────────────────────────────────
# Valores de referencia (archivo 20261026, moneda de reporte de cada marca)
# ──────────────────────────────────────────────────────────────────────

# marca, moneda, fact, bo, forecast, cumpl anual (1 dec), trimestres (enteros; None = sin forecast)
CALENDARIO_CON_FORECAST = [
    ("Steelco", "EUR", 1502302, 221389.69, 1250000, 137.9, [86, 106, 200, 160]),
    ("Melag", "EUR", 276689, 14238, 520200, 55.9, [146, 142, 31, 29]),
    ("Pentax", "USD", 758351, 35631.22, 1249999.2, 63.5, [97, 60, 87, 9]),
    ("Eurosets", "EUR", 656301, 582510.16, 800000, 154.9, [72, 118, 140, 290]),
    ("Creo Medical", "USD", 59776, 32999.6, 90000, 103.1, [42, None, None, None]),
    ("Bien Air", "USD", 58514.4, 0, 50400, 116.1, [104, 84, 250, 27]),
]

CALENDARIO_SIN_FORECAST = [
    ("DDC Dolphin", "GBP", 17571, 35232.91),
    ("Yipak", "USD", 147402, 0),
    ("Franklab", "EUR", 3468, 0),
    ("G-Flex", "EUR", 60950.8, 18260.28),
    ("Durr Dental", "EUR", 86862, 0),
    ("Normmed", "EUR", 34522.1, 0),
]

FISCAL = [
    ("Pentax", 1211512, 3710, 1249999.2, 97, [63, 169, 97, 60]),
    ("G-Flex", 149250.4, 0, None, None, [None, None, None, None]),
    ("Durr Dental", 41870, 0, None, None, [None, None, None, None]),
]


@pytest.mark.parametrize("nombre,moneda,fact,bo,fc,cumpl,trims", CALENDARIO_CON_FORECAST)
def test_calendario_con_forecast(datos, nombre, moneda, fact, bo, fc, cumpl, trims):
    c = bloque(datos, nombre)
    assert marca(datos, nombre)["moneda"] == moneda
    assert c["total"]["fact"] == pytest.approx(fact, abs=0.01)
    assert c["total"]["bo"] == pytest.approx(bo, abs=0.01)
    assert c["total"]["fc"] == pytest.approx(fc, abs=0.01)
    assert pct(c["total"]["cumpl"], 1) == cumpl
    assert [pct(t["cumpl"]) for t in c["trims"]] == trims


def test_bien_air_trimestres_con_decimal(datos):
    c = bloque(datos, "Bien Air")
    assert pct(c["total"]["cumpl"], 1) == 116.1
    assert [pct(t["cumpl"], 1) for t in c["trims"]] == [104.3, 83.5, 249.6, 27.1]


@pytest.mark.parametrize("nombre,moneda,fact,bo", CALENDARIO_SIN_FORECAST)
def test_calendario_sin_forecast(datos, nombre, moneda, fact, bo):
    c = bloque(datos, nombre)
    assert marca(datos, nombre)["moneda"] == moneda
    assert c["total"]["fact"] == pytest.approx(fact, abs=0.01)
    assert c["total"]["bo"] == pytest.approx(bo, abs=0.01)
    assert c["tieneForecast"] is False
    # Sin forecast = "—" (None), nunca 0 %
    assert c["total"]["cumpl"] is None
    assert c["ytd"]["cumpl"] is None
    assert all(t["cumpl"] is None for t in c["trims"])


@pytest.mark.parametrize("nombre,fact,bo,fc,cumpl,trims", FISCAL)
def test_bloque_fiscal(datos, nombre, fact, bo, fc, cumpl, trims):
    m = marca(datos, nombre)
    assert m["tieneFiscal"] is True
    b = next(x for x in m["bloques"] if x["tipo"] == "fiscal")
    assert b["titulo"] == "Año fiscal USA (Jul-25 – Jun-26)"
    c = b["calc"]
    assert c["total"]["fact"] == pytest.approx(fact, abs=0.01)
    assert c["total"]["bo"] == pytest.approx(bo, abs=0.01)
    if fc is None:
        assert c["tieneForecast"] is False
        assert c["total"]["cumpl"] is None
    else:
        assert c["total"]["fc"] == pytest.approx(fc, abs=0.01)
        assert pct(c["total"]["cumpl"]) == cumpl
    assert [pct(t["cumpl"]) for t in c["trims"]] == trims


def test_columna_etiquetas_fiscal(datos):
    # U en Pentax y Durr Dental, V en G-Flex
    col = {m["nombre"]: gd.col_letra(b["c0"]) for m in datos["marcas"] for b in m["bloques"] if b["tipo"] == "fiscal"}
    assert col == {"Pentax": "U", "G-Flex": "V", "Durr Dental": "U"}


def test_conciliacion_cuadra(datos):
    r = datos["resumen"]
    assert r["comparaciones"] == 165
    assert r["diferencias"] == 0
    assert r["erroresBloqueantes"] == 0
    assert all(i["ok"] for m in datos["marcas"] for b in m["bloques"] for i in b["conc"])


def test_lee_hoja_exacta_e_ignora_la_del_espacio(datos):
    assert datos["meta"]["fact"]["hoja"] == "BBDD Facturacion"
    assert datos["meta"]["fact"]["filas"] == 699
    assert datos["meta"]["bo"]["filas"] == 1477
    assert any("BBDD␣Facturacion␣" in p["titulo"] for p in datos["problemas"])


def test_solo_hojas_de_marca_visibles(datos):
    assert [m["nombre"] for m in datos["marcas"]] == [
        "Steelco", "Melag", "DDC Dolphin", "Yipak", "Franklab", "Pentax", "Creo Medical",
        "G-Flex", "Bien Air", "Durr Dental", "Eurosets", "Normmed"]
    # Clave de marca = A1 con trim, no el nombre de la hoja
    assert marca(datos, "Creo Medical")["hoja"] == "Creomedical"
    assert marca(datos, "Bien Air")["hoja"] == "Bien Air "


def test_cortes(datos):
    assert datos["meta"]["fact"]["corte"] == "2026-10-20"
    assert datos["meta"]["bo"]["corte"] == "2026-10-26"


def test_ytd_es_enero_a_mes_de_corte(datos):
    c = bloque(datos, "Steelco")
    y = c["ytd"]
    assert y["nMeses"] == 10 and y["hasta"] == {"anio": 2026, "mes": 9}
    assert y["fc"] == pytest.approx(sum(m["fc"] for m in c["meses"][:10]))
    assert y["factbo"] == pytest.approx(sum(m["factbo"] for m in c["meses"][:10]))
    assert y["cumpl"] == pytest.approx(y["factbo"] / y["fc"])
    assert [t["estado"] for t in c["trims"]] == ["cerrado", "cerrado", "cerrado", "parcial"]


def test_conversion_de_moneda(datos):
    assert {m["nombre"] for m in datos["marcas"] if m["convierte"]} == {"G-Flex", "Bien Air", "Normmed"}
    pares = {(p["orig"], p["dest"], p["tc"]) for p in marca(datos, "G-Flex")["paresConversion"]}
    assert pares == {("USD", "EUR", 0.9)}


def test_sin_detalle_no_incluye_oc_productos_ni_precios():
    ruta = _excel_real()
    if ruta is None:
        pytest.skip("No hay Excel en data/raw/.")
    d = gd.procesar_libro(ruta, incluir_detalle=False)
    assert d["meta"]["incluirDetalle"] is False
    assert all("detalle" not in m for m in d["marcas"])
    texto = json.dumps(d, ensure_ascii=False, default=str)
    for campo in ('"referencia"', '"producto"', '"precio"', '"folio"', '"oc"', '"descripcion"'):
        assert campo not in texto


# ──────────────────────────────────────────────────────────────────────
# Reglas de negocio con un libro sintético (no requiere el Excel real)
# ──────────────────────────────────────────────────────────────────────

MESES_HOJA = ["Enero", "Febrero", "Marzo", " Q1 ", "Abril", "Mayo", "Junio", " Q2 ", "Julio", "Agosto",
              "Septiembre", " Q3 ", "Octubre", "Noviembre", "Diciembre", " Q4 ", " Total "]


def _hoja_marca(wb, titulo, a1, forecast=None, oculta=False):
    ws = wb.create_sheet(titulo)
    ws["A1"], ws["A2"], ws["A6"] = a1, " Item ", "Forecast "
    for i, etiqueta in enumerate(MESES_HOJA, start=2):
        ws.cell(row=1, column=i, value=2026)
        ws.cell(row=2, column=i, value=etiqueta)
    for col, v in (forecast or {}).items():
        ws[f"{col}6"] = v
    if oculta:
        ws.sheet_state = "hidden"
    return ws


def libro_sintetico(tmp_path, hoja_fact="BBDD Facturacion", sin_encabezado=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = hoja_fact
    enc = [h for h in gd.ENCABEZADOS_FACTURACION if h != sin_encabezado]
    ws.append(enc)
    fila = {"Proveedor": "Marca X ", "Fecha de contabilización": datetime(2026, 2, 10), "Año": 2026, "Mes": "febrero",
            "Número de folio": 1, "Comentarios": "", "Orden de compra": "", "Categoria": "Equipos",
            "Moneda del documento": "EUR", "Total documento (ME)": 100, "Moneda marca": "EUR",
            "TC a moneda marca": 1, "Total moneda marca": 100}
    ws.append([fila[h] for h in enc])
    # Una fila de otra marca y una con categoría no estándar
    ws.append([{**fila, "Proveedor": "Otra", "Total moneda marca": 999}[h] for h in enc])
    ws.append([{**fila, "Categoria": 0, "Total moneda marca": 7}[h] for h in enc])

    bo = wb.create_sheet("BBDD BO abiertas")
    bo.append(gd.ENCABEZADOS_BO)
    fila_bo = {"Fecha OC": datetime(2026, 5, 5), "Año OC": 2026, "Mes": "mayo", "Referencia": "0001/26", "Producto": "P1",
               "Descripcion": "Prod", "Grupo": "ERROR", "Marca": "Marca X", "Cantidad pdte": 2, "moneda": "EUR",
               "Precio": 25, "Total Neto ME": 50, "Moneda marca": "EUR", "TC a moneda marca": 1, "Total moneda marca": 50}
    bo.append([fila_bo[h] for h in gd.ENCABEZADOS_BO])

    pm = wb.create_sheet("Parametros Moneda")
    pm.append(["Marca", "Moneda reporte", None, "Par", "Tipo de cambio"])

    # Forecast solo trimestral en Q1 (E6) → meses derivados = 300/3
    _hoja_marca(wb, "MarcaX", "Marca X", forecast={"E": 300, "R": 300})
    _hoja_marca(wb, "Oculta", "Otra", forecast={"E": 1, "R": 1}, oculta=True)
    ruta = tmp_path / "sintetico.xlsx"
    wb.save(ruta)
    return ruta


def test_sintetico_reglas(tmp_path):
    d = gd.procesar_libro(libro_sintetico(tmp_path), incluir_detalle=True)
    assert [m["nombre"] for m in d["marcas"]] == ["Marca X"]  # la hoja oculta no se lee
    c = d["marcas"][0]["bloques"][0]["calc"]
    # Proveedor con espacio final se compara con trim
    assert c["total"]["fact"] == 107
    assert c["total"]["bo"] == 50
    # Forecast derivado: trimestre / 3 cuando los meses están vacíos
    assert [m["fc"] for m in c["meses"][:3]] == [100, 100, 100]
    assert all(m["fcDerivado"] for m in c["meses"][:3])
    assert c["meses"][3]["fc"] is None
    # Categorías: 0 y "ERROR" van a "Sin categoría"
    assert c["catF"]["Equipos"] == 100 and c["catF"]["Sin categoría"] == 7
    assert c["catB"]["Sin categoría"] == 50
    # Cumplimiento: Q1 con forecast; Q2–Q4 sin forecast → None
    assert c["trims"][0]["cumpl"] == pytest.approx(107 / 300)
    assert [t["cumpl"] for t in c["trims"][1:]] == [None, None, None]
    assert c["total"]["cumpl"] == pytest.approx(157 / 300)
    # Corte = 05-05-2026 → YTD ene–may, Q2 parcial, Q3 y Q4 futuros
    assert c["ytd"]["nMeses"] == 5
    assert c["ytd"]["cumpl"] == pytest.approx(157 / 300)
    assert [t["estado"] for t in c["trims"]] == ["cerrado", "parcial", "futuro", "futuro"]


def test_sintetico_solo_hoja_con_espacio_falla(tmp_path):
    ruta = libro_sintetico(tmp_path, hoja_fact="BBDD Facturacion ")
    with pytest.raises(gd.ErrorCarga, match="nombre exacto"):
        gd.procesar_libro(ruta)


def test_sintetico_encabezado_faltante_falla(tmp_path):
    ruta = libro_sintetico(tmp_path, sin_encabezado="Total moneda marca")
    with pytest.raises(gd.ErrorCarga, match='"Total moneda marca"'):
        gd.procesar_libro(ruta)


def test_sintetico_conciliacion_detecta_diferencias(tmp_path):
    # Las celdas de la hoja sintética están vacías (Excel = 0) y la app suma 107 → debe marcar diferencia.
    d = gd.procesar_libro(libro_sintetico(tmp_path))
    assert d["resumen"]["diferencias"] > 0


# ──────────────────────────────────────────────────────────────────────
# JSON publicado
# ──────────────────────────────────────────────────────────────────────

def test_json_publicado_es_valido_y_conciliado():
    ruta = RAIZ / gd.SALIDA_JSON
    if not ruta.exists():
        pytest.skip("Aún no se generó docs/datos_cumplimiento.json.")
    d = json.loads(ruta.read_text(encoding="utf-8"))
    assert d["resumen"]["diferencias"] == 0
    assert d["meta"]["fact"]["hoja"] == "BBDD Facturacion"
    assert d["meta"]["archivo"].endswith(".xlsx")
    assert ("detalle" in d["marcas"][0]) == d["meta"]["incluirDetalle"]
